import io
import json

import pytest
from openpyxl import Workbook

from app.structured_intake import TableIntakeError, read_table_bytes


def workbook_bytes(rows, extra_sheet=False):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    if extra_sheet:
        book.create_sheet("Other")
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def test_json_and_xlsx_match_csv_preserving_postal_and_identity():
    rows = [{"PAT_ID": "001", "ZIP": "00901", "AGE": 43, "DOS": "2026-10-17"},
            {"PAT_ID": "002", "ZIP": "02101", "AGE": 52, "DOS": "2026-10-18"}]
    reference = read_table_bytes(b"PAT_ID,ZIP,AGE,DOS\n001,00901,43,2026-10-17\n002,02101,52,2026-10-18\n", "data.csv")
    json_frame = read_table_bytes(json.dumps(rows).encode(), "data.json")
    xlsx_frame = read_table_bytes(workbook_bytes([list(rows[0]), *[list(r.values()) for r in rows]]), "data.xlsx")
    assert reference.equals(json_frame)
    assert reference.equals(xlsx_frame)
    assert reference.ZIP.tolist() == ["00901", "02101"]
    assert reference.PAT_ID.tolist() == ["001", "002"]


@pytest.mark.parametrize("payload", [b'[{"x": [1,2]}]', b'[{"x":1},{"y":2}]', b'[{"x":NaN}]', b'[]'])
def test_json_rejects_nonrectangular_or_nonfinite_tables(payload):
    with pytest.raises(TableIntakeError):
        read_table_bytes(payload, "data.json")


def test_xlsx_does_not_silently_use_cached_formulas_or_ignore_other_sheets():
    with pytest.raises(TableIntakeError, match="formulas"):
        read_table_bytes(workbook_bytes([["AGE"], ["=40+3"]]), "formula.xlsx")
    with pytest.raises(TableIntakeError, match="one visible"):
        read_table_bytes(workbook_bytes([["AGE"], [43]], extra_sheet=True), "multiple.xlsx")


def test_json_records_wrapper_and_nulls():
    result = read_table_bytes(b'{"records":[{"age":43,"zip":"00901"},{"age":null,"zip":"02101"}]}', "table.json")
    assert len(result) == 2
    assert result.age.isna().sum() == 1


def test_duplicate_json_keys_are_not_silently_overwritten():
    with pytest.raises(TableIntakeError, match="distinct"):
        read_table_bytes(b'[{"age":43,"age":44}]', "table.json")


def test_xlsx_preserves_explicit_zero_padded_postal_codes():
    book = Workbook()
    book.active.append(["ZIP", "AGE"])
    book.active.append([901, 43])
    book.active["A2"].number_format = "00000"
    output = io.BytesIO()
    book.save(output)
    result = read_table_bytes(output.getvalue(), "table.xlsx")
    assert result.ZIP.iloc[0] == "00901"
