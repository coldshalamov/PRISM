"""Bounded table intake. All formats share the existing CSV normalization path."""
from __future__ import annotations

import csv
import io
import json
import math
import re
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from . import datasets


class TableIntakeError(ValueError):
    """Messages contain only fixed format guidance, never source values."""


def _json_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise TableIntakeError("JSON object keys must be distinct; duplicate columns cannot be imported.")
        result[key] = value
    return result


def _excel_value(cell):
    value = cell.value
    # Preserve explicit zero-padded codes without interpreting arbitrary Excel
    # formatting, locale, currency or formulas as data transformations.
    number_format = getattr(cell, "number_format", "") or ""
    if (isinstance(value, int) and not isinstance(value, bool) and value >= 0
            and re.fullmatch(r"0{2,20}", number_format)):
        return str(value).zfill(len(number_format))
    return value


def _table(headers, rows):
    if (not headers or len(headers) > 80 or any(not isinstance(h, str) for h in headers)
            or len(set(headers)) != len(headers)):
        raise TableIntakeError("Use 1–80 distinct text column headings.")
    target = io.StringIO(newline="")
    writer = csv.writer(target)
    writer.writerow(headers)
    count = 0
    for row in rows:
        count += 1
        if count > datasets.MAX_ROWS:
            raise TableIntakeError("Tables can contain at most 100,000 data rows.")
        if len(row) != len(headers):
            raise TableIntakeError("Every row must match the table's column headings.")
        values = []
        for value in row:
            if isinstance(value, (dict, list, tuple)):
                raise TableIntakeError("Nested objects and arrays are not table cells. Flatten them before import.")
            if isinstance(value, float) and not math.isfinite(value):
                raise TableIntakeError("Table cells cannot contain infinity or NaN literals.")
            if isinstance(value, (date, datetime)):
                value = value.isoformat()
            if value is not None and len(str(value)) > datasets.MAX_CELL_CHARS:
                raise TableIntakeError("Table cells can contain at most 65,536 characters.")
            values.append(value)
        writer.writerow(values)
        if target.tell() > datasets.MAX_CSV_BYTES:
            raise TableIntakeError("The normalized table exceeds the 50 MB limit.")
    if not count:
        raise TableIntakeError("The table must contain at least one data row.")
    return datasets.read_csv_bytes(target.getvalue().encode("utf-8"))


def read_table_bytes(content: bytes, filename: str):
    if not content or len(content) > datasets.MAX_CSV_BYTES:
        raise TableIntakeError("Choose a nonempty table of at most 50 MB.")
    suffix = Path(filename or "table.csv").suffix.lower()
    if suffix == ".csv":
        return datasets.read_csv_bytes(content)
    if suffix == ".json":
        try:
            records = json.loads(content.decode("utf-8-sig"), object_pairs_hook=_json_object,
                                 parse_constant=lambda _: (_ for _ in ()).throw(ValueError()))
        except TableIntakeError:
            raise
        except (UnicodeDecodeError, ValueError, RecursionError):
            raise TableIntakeError("Use valid UTF-8 JSON containing an array of table records.") from None
        if isinstance(records, dict) and set(records) == {"records"}:
            records = records["records"]
        if not isinstance(records, list) or not records or not isinstance(records[0], dict):
            raise TableIntakeError("JSON must be a nonempty array of objects, or an object with a records array.")
        headers = list(records[0])
        if any(not isinstance(row, dict) or set(row) != set(headers) for row in records):
            raise TableIntakeError("JSON records must have the same column names; use null for missing cells.")
        return _table(headers, ([row[h] for h in headers] for row in records))
    if suffix == ".xlsx":
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                    raise TableIntakeError("The expanded workbook exceeds the 200 MB limit.")
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        except TableIntakeError:
            raise
        except (BadZipFile, ValueError, KeyError, OSError):
            raise TableIntakeError("The XLSX workbook could not be read. Export a single table as CSV or JSON.") from None
        try:
            if len(workbook.worksheets) != 1 or workbook.worksheets[0].sheet_state != "visible":
                raise TableIntakeError("Use a workbook with exactly one visible worksheet. Export other sheets separately.")
            sheet = workbook.worksheets[0]
            if sheet.max_column and sheet.max_column > 80:
                raise TableIntakeError("The worksheet can contain at most 80 columns.")
            if sheet.max_row and sheet.max_row > datasets.MAX_ROWS + 1:
                raise TableIntakeError("The worksheet can contain at most 100,000 data rows.")
            iterator = sheet.iter_rows()
            header_cells = next(iterator, ())
            headers = [cell.value for cell in header_cells]
            if any(cell.data_type == "f" for cell in header_cells):
                raise TableIntakeError("Replace spreadsheet formulas with explicit values before import.")

            def rows():
                for index, row in enumerate(iterator):
                    if index >= datasets.MAX_ROWS:
                        raise TableIntakeError("The worksheet can contain at most 100,000 data rows.")
                    if any(cell.data_type == "f" for cell in row):
                        raise TableIntakeError("Replace spreadsheet formulas with explicit values before import.")
                    values = [_excel_value(cell) for cell in row]
                    if any(value is not None for value in values):
                        yield values

            return _table(headers, rows())
        finally:
            workbook.close()
    raise TableIntakeError("Choose a CSV, JSON table, or single-sheet XLSX workbook.")
